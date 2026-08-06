"""
Pure-Python SocialCalc <-> Excel/CSV export and import.

Replaces the PHPExcel-based PHP scripts (export.php / import.php) which are
incompatible with PHP 8+.
"""

import csv
import io
import json
import re
import tempfile
from collections import OrderedDict

import openpyxl
import xlwt
import xlrd


def _coord_to_cr(coord):
    """Convert 'A1' -> (col, row) 1-based."""
    m = re.match(r'^([A-Z]+)(\d+)$', coord.upper())
    if not m:
        return (1, 1)
    letters, digits = m.group(1), int(m.group(2))
    col = 0
    for ch in letters:
        col = col * 26 + (ord(ch) - ord('A') + 1)
    return (col, digits)


def _decode_value(val):
    """Decode SocialCalc encoded values (backslash escapes)."""
    if val is None:
        return ''
    val = val.replace('\\c', ':')
    val = val.replace('\\n', '\n')
    val = val.replace('\\b', '\\')
    return val


def _parse_sheet(data):
    """Parse a SocialCalc sheet savestr into a dict of cells and metadata."""
    cells = {}
    for line in data.split('\n'):
        line = line.strip()
        if not line:
            continue
        parts = line.split(':')
        if parts[0] == 'cell' and len(parts) >= 3:
            coord = parts[1].upper()
            cell = _parse_cell(parts, 2)
            cells[coord] = cell
    return cells


def _parse_cell(parts, j):
    """Parse cell properties from colon-separated parts."""
    cell = {}
    while j < len(parts):
        t = parts[j]
        j += 1
        if t == 'v':
            try:
                cell['datavalue'] = float(_decode_value(parts[j]))
            except (ValueError, IndexError):
                cell['datavalue'] = 0
            cell['datatype'] = 'v'
            cell['valuetype'] = 'n'
            j += 1
        elif t == 't':
            cell['datavalue'] = _decode_value(parts[j]) if j < len(parts) else ''
            cell['datatype'] = 't'
            cell['valuetype'] = 't'
            j += 1
        elif t == 'vt':
            if j + 1 < len(parts):
                vtype = parts[j]
                cell['valuetype'] = vtype
                cell['datatype'] = 'v' if vtype.startswith('n') else 't'
                cell['datavalue'] = _decode_value(parts[j + 1])
                if cell['datatype'] == 'v':
                    try:
                        cell['datavalue'] = float(cell['datavalue'])
                    except (ValueError, TypeError):
                        pass
                j += 2
        elif t == 'vtf':
            if j + 2 < len(parts):
                cell['valuetype'] = parts[j]
                cell['datavalue'] = _decode_value(parts[j + 1])
                cell['formula'] = _decode_value(parts[j + 2])
                cell['datatype'] = 'f'
                if cell['valuetype'].startswith('n'):
                    try:
                        cell['datavalue'] = float(cell['datavalue'])
                    except (ValueError, TypeError):
                        pass
                j += 3
        elif t == 'vtc':
            if j + 2 < len(parts):
                cell['valuetype'] = parts[j]
                cell['datavalue'] = _decode_value(parts[j + 1])
                cell['formula'] = _decode_value(parts[j + 2])
                cell['datatype'] = 'c'
                if cell['valuetype'].startswith('n'):
                    try:
                        cell['datavalue'] = float(cell['datavalue'])
                    except (ValueError, TypeError):
                        pass
                j += 3
        elif t == 'f':
            if j < len(parts):
                cell['font'] = parts[j]
                j += 1
        elif t == 'b':
            j += 4
        elif t == 'l':
            j += 1
        elif t == 'c':
            if j < len(parts):
                cell['color'] = parts[j]
                j += 1
        elif t == 'bg':
            j += 1
        elif t == 'cf':
            j += 1
        elif t == 'ntvf':
            j += 1
        elif t == 'tvf':
            j += 1
        elif t == 'colspan':
            if j < len(parts):
                cell['colspan'] = int(parts[j])
                j += 1
        elif t == 'rowspan':
            if j < len(parts):
                cell['rowspan'] = int(parts[j])
                j += 1
        elif t == 'cssc':
            j += 1
        elif t == 'csss':
            j += 1
        elif t == 'mod':
            j += 1
        elif t == 'comment':
            if j < len(parts):
                cell['comment'] = _decode_value(parts[j])
                j += 1
        elif t == 'ro':
            j += 1
        elif t == 'e':
            j += 1
        else:
            pass
    return cell


def _write_cells_to_xlsx_sheet(ws, cells):
    """Write parsed cells to an openpyxl worksheet."""
    for coord, cell in cells.items():
        col, row = _coord_to_cr(coord)
        val = cell.get('datavalue', '')
        formula = cell.get('formula')
        if formula and cell.get('datatype') == 'f':
            ws.cell(row=row, column=col, value='=' + formula)
        else:
            ws.cell(row=row, column=col, value=val)
        if cell.get('colspan') or cell.get('rowspan'):
            colspan = cell.get('colspan', 1)
            rowspan = cell.get('rowspan', 1)
            if colspan > 1 or rowspan > 1:
                ws.merge_cells(
                    start_row=row, start_column=col,
                    end_row=row + rowspan - 1, end_column=col + colspan - 1
                )


def _write_cells_to_xls_sheet(ws, cells):
    """Write parsed cells to an xlwt worksheet."""
    for coord, cell in cells.items():
        col, row = _coord_to_cr(coord)
        val = cell.get('datavalue', '')
        formula = cell.get('formula')
        if formula and cell.get('datatype') == 'f':
            try:
                ws.write(row - 1, col - 1, xlwt.Formula(formula))
            except Exception:
                ws.write(row - 1, col - 1, val)
        else:
            ws.write(row - 1, col - 1, val)


def export_xlsx(json_data):
    """Export SocialCalc JSON to .xlsx bytes."""
    import logging as _log
    _log.info("[DEBUG-INTEROP export_xlsx] input type=%s len=%d first_300=%s", type(json_data).__name__, len(json_data) if json_data else 0, repr(json_data[:300]) if json_data else 'None')
    book = json.loads(json_data) if isinstance(json_data, str) else json_data
    _log.info("[DEBUG-INTEROP export_xlsx] book keys=%s sheetArr keys=%s", list(book.keys()) if isinstance(book, dict) else 'NOT_DICT', list(book.get('sheetArr', {}).keys()) if isinstance(book, dict) else 'N/A')
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    sheet_arr = book.get('sheetArr', {})
    for idx, (key, sheet_info) in enumerate(sheet_arr.items()):
        title = sheet_info.get('name', f'Sheet{idx + 1}')
        ws = wb.create_sheet(title=title[:31])
        savestr = sheet_info.get('sheetstr', {}).get('savestr', '')
        cells = _parse_sheet(savestr)
        _write_cells_to_xlsx_sheet(ws, cells)

    if not wb.sheetnames:
        wb.create_sheet('Sheet1')

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_xls(json_data):
    """Export SocialCalc JSON to .xls bytes."""
    book = json.loads(json_data) if isinstance(json_data, str) else json_data
    wb = xlwt.Workbook(encoding='utf-8')

    sheet_arr = book.get('sheetArr', {})
    for idx, (key, sheet_info) in enumerate(sheet_arr.items()):
        title = sheet_info.get('name', f'Sheet{idx + 1}')
        ws = wb.add_sheet(title[:31])
        savestr = sheet_info.get('sheetstr', {}).get('savestr', '')
        cells = _parse_sheet(savestr)
        _write_cells_to_xls_sheet(ws, cells)

    if not sheet_arr:
        wb.add_sheet('Sheet1')

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_csv(json_data):
    """Export SocialCalc JSON to CSV string (first sheet only)."""
    import logging as _log
    _log.info("[DEBUG-INTEROP export_csv] input type=%s len=%d first_300=%s", type(json_data).__name__, len(json_data) if json_data else 0, repr(json_data[:300]) if json_data else 'None')
    book = json.loads(json_data) if isinstance(json_data, str) else json_data
    _log.info("[DEBUG-INTEROP export_csv] book type=%s keys=%s", type(book).__name__, list(book.keys()) if isinstance(book, dict) else repr(book)[:200])
    sheet_arr = book.get('sheetArr', {})
    _log.info("[DEBUG-INTEROP export_csv] sheet_arr type=%s len=%d", type(sheet_arr).__name__, len(sheet_arr) if sheet_arr else 0)

    if not sheet_arr:
        return ''

    first_sheet = next(iter(sheet_arr.values()))
    savestr = first_sheet.get('sheetstr', {}).get('savestr', '')
    cells = _parse_sheet(savestr)

    if not cells:
        return ''

    max_row = 0
    max_col = 0
    for coord in cells:
        col, row = _coord_to_cr(coord)
        max_row = max(max_row, row)
        max_col = max(max_col, col)

    output = io.StringIO()
    writer = csv.writer(output)
    for r in range(1, max_row + 1):
        row_data = []
        for c in range(1, max_col + 1):
            col_letter = ''
            temp_c = c
            while temp_c > 0:
                temp_c, remainder = divmod(temp_c - 1, 26)
                col_letter = chr(65 + remainder) + col_letter
            coord = f'{col_letter}{r}'
            cell = cells.get(coord, {})
            val = cell.get('datavalue', '')
            row_data.append(val)
        writer.writerow(row_data)

    return output.getvalue()


def _col_letter(c):
    """Convert 1-based column number to Excel letter(s)."""
    result = ''
    while c > 0:
        c, remainder = divmod(c - 1, 26)
        result = chr(65 + remainder) + result
    return result


def _encode_value(val):
    """Encode a value for SocialCalc format."""
    if val is None:
        return ''
    s = str(val)
    s = s.replace('\\', '\\b')
    s = s.replace(':', '\\c')
    s = s.replace('\n', '\\n')
    return s


def import_xlsx(file_bytes):
    """Import .xlsx file bytes and return SocialCalc JSON structure."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    return _workbook_to_socialcalc_openpyxl(wb)


def import_xls(file_bytes):
    """Import .xls file bytes and return SocialCalc JSON structure."""
    with tempfile.NamedTemporaryFile(suffix='.xls', delete=False) as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
    try:
        wb = xlrd.open_workbook(tmp_path)
        return _workbook_to_socialcalc_xlrd(wb)
    finally:
        import os as _os
        _os.unlink(tmp_path)


def import_csv(file_bytes):
    """Import CSV file bytes and return SocialCalc JSON structure."""
    text = file_bytes.decode('utf-8', errors='replace')
    reader = csv.reader(io.StringIO(text))

    lines = []
    for r, row in enumerate(reader, 1):
        for c, val in enumerate(row, 1):
            if val == '':
                continue
            coord = f'{_col_letter(c)}{r}'
            try:
                numval = float(val)
                lines.append(f'cell:{coord}:v:{_encode_value(val)}')
            except ValueError:
                lines.append(f'cell:{coord}:t:{_encode_value(val)}')
    savestr = 'version:1.5\n' + '\n'.join(lines) + '\n'

    book = {
        'numsheets': 1,
        'currentname': 'Sheet1',
        'currentid': 'sheet0',
        'sheetArr': {
            'Sheet0': {
                'name': 'Sheet1',
                'sheetstr': {'savestr': savestr}
            }
        }
    }
    return json.dumps(book)


def _workbook_to_socialcalc_openpyxl(wb):
    """Convert openpyxl workbook to SocialCalc JSON."""
    book = {
        'numsheets': len(wb.sheetnames),
        'currentname': wb.active.title if wb.active else wb.sheetnames[0],
        'sheetArr': {}
    }

    for idx, name in enumerate(wb.sheetnames):
        ws = wb[name]
        lines = ['version:1.5']
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                coord = cell.coordinate
                val = cell.value
                if isinstance(val, (int, float)):
                    lines.append(f'cell:{coord}:v:{_encode_value(val)}')
                else:
                    lines.append(f'cell:{coord}:t:{_encode_value(val)}')
        savestr = '\n'.join(lines) + '\n'
        sheet_key = f'Sheet{idx}'
        book['sheetArr'][sheet_key] = {
            'name': name,
            'sheetstr': {'savestr': savestr}
        }
        if name == book['currentname']:
            book['currentid'] = f'sheet{idx}'

    if 'currentid' not in book and book['numsheets'] > 0:
        book['currentid'] = 'sheet0'

    return json.dumps(book)


def _workbook_to_socialcalc_xlrd(wb):
    """Convert xlrd workbook to SocialCalc JSON."""
    book = {
        'numsheets': wb.nsheets,
        'currentname': wb.sheet_names()[0] if wb.nsheets > 0 else 'Sheet1',
        'sheetArr': {}
    }

    for idx in range(wb.nsheets):
        ws = wb.sheet_by_index(idx)
        name = ws.name
        lines = ['version:1.5']
        for row in range(ws.nrows):
            for col in range(ws.ncols):
                cell = ws.cell(row, col)
                if cell.value == '' or cell.value is None:
                    continue
                coord = f'{_col_letter(col + 1)}{row + 1}'
                if cell.ctype in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_DATE):
                    lines.append(f'cell:{coord}:v:{_encode_value(cell.value)}')
                else:
                    lines.append(f'cell:{coord}:t:{_encode_value(cell.value)}')
        savestr = '\n'.join(lines) + '\n'
        sheet_key = f'Sheet{idx}'
        book['sheetArr'][sheet_key] = {
            'name': name,
            'sheetstr': {'savestr': savestr}
        }
        if name == book['currentname']:
            book['currentid'] = f'sheet{idx}'

    if 'currentid' not in book and book['numsheets'] > 0:
        book['currentid'] = 'sheet0'

    return json.dumps(book)
